import os
import docx
from docx import Document

# Chemin vers le dossier des templates docx
TEMPLATE_DIR = os.path.join("templates", "docx")
os.makedirs(TEMPLATE_DIR, exist_ok=True)


def generer_template_service():
    doc = Document()

    doc.add_heading("FACTURE DE SERVICE", level=1)

    # En-tête / Informations client
    p = doc.add_paragraph()
    p.add_run("Date : ").bold = True
    p.add_run("{{ date_facture }}\n")
    p.add_run("N° Facture : ").bold = True
    p.add_run("{{ numero_facture }}\n")
    p.add_run("Client : ").bold = True
    p.add_run("{{ client_nom }}\n")
    p.add_run("NIF : ").bold = True
    p.add_run("{{ client_nif }}\n")
    p.add_run("BP : ").bold = True
    p.add_run("{{ client_bp }}\n")
    p.add_run("Ville : ").bold = True
    p.add_run("{{ client_ville }}\n")

    doc.add_paragraph("Objet : {{ objet }}")

    # Tableau des prestations
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    # En-tête du tableau
    hdr = table.rows[0].cells
    hdr[0].text = "Désignation"
    hdr[1].text = "Nombre de jours / Période"
    hdr[2].text = "Prix journalier / Taux"
    hdr[3].text = "Prix total"

    # Ligne dynamique Jinja2
    row = table.add_row().cells
    row[0].text = "{% for item in prestations %}{{ item.designation }}"
    row[1].text = "{{ item.periode }}"
    row[2].text = "{{ item.taux }}"
    row[3].text = "{{ item.total }}{% endfor %}"

    # Totaux et options
    r_ttc = table.add_row().cells
    r_ttc[0].text = "MONTANT TTC"
    r_ttc[3].text = "{{ montant_ttc }}"

    r_prelev = table.add_row().cells
    r_prelev[
        0
    ].text = "{% if has_prelevement %}Prélèvement ({{ taux_prelevement }}%){% endif %}"
    r_prelev[3].text = "{% if has_prelevement %}{{ prelevement }}{% endif %}"

    r_av = table.add_row().cells
    r_av[0].text = "{% if has_avance %}Avance{% endif %}"
    r_av[3].text = "{% if has_avance %}{{ avance }}{% endif %}"

    r_reg = table.add_row().cells
    r_reg[0].text = "{% if has_regule %}Régule{% endif %}"
    r_reg[3].text = "{% if has_regule %}{{ regule }}{% endif %}"

    r_opt = table.add_row().cells
    r_opt[0].text = "{% if has_autre_option %}{{ autre_option_label }}{% endif %}"
    r_opt[3].text = (
        "{% if has_autre_option %}{{ autre_option_montant }}{% endif %}"
    )

    r_net = table.add_row().cells
    r_net[0].text = "NET À PAYER"
    r_net[3].text = "{{ net_a_payer }}"

    doc.add_paragraph(
        "\nArrêté la présente facture à la somme de : {{ net_en_lettres }}"
    )

    filepath = os.path.join(TEMPLATE_DIR, "Facture_Atelier_Mobile.docx")
    doc.save(filepath)
    print(f" Template réinitialisé : {filepath}")


def generer_template_produit():
    doc = Document()

    doc.add_heading("FACTURE DE VENTE PRODUITS", level=1)

    # En-tête / Informations client
    p = doc.add_paragraph()
    p.add_run("Date : ").bold = True
    p.add_run("{{ date_facture }}\n")
    p.add_run("N° Facture : ").bold = True
    p.add_run("{{ numero_facture }}\n")
    p.add_run("Client : ").bold = True
    p.add_run("{{ client_nom }}\n")
    p.add_run("NIF : ").bold = True
    p.add_run("{{ client_nif }}\n")
    p.add_run("BP : ").bold = True
    p.add_run("{{ client_bp }}\n")
    p.add_run("Ville : ").bold = True
    p.add_run("{{ client_ville }}\n")

    doc.add_paragraph("Objet : {{ objet }}")

    # Tableau des produits
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    # En-tête du tableau
    hdr = table.rows[0].cells
    hdr[0].text = "Désignation"
    hdr[1].text = "Quantité"
    hdr[2].text = "Prix Unitaire"
    hdr[3].text = "Total HT"

    # Ligne dynamique Jinja2
    row = table.add_row().cells
    row[0].text = "{% for item in prestations %}{{ item.designation }}"
    row[1].text = "{{ item.quantite }}"
    row[2].text = "{{ item.prix_unitaire }}"
    row[3].text = "{{ item.total }}{% endfor %}"

    # Totaux et options
    r_ttc = table.add_row().cells
    r_ttc[0].text = "MONTANT TTC"
    r_ttc[3].text = "{{ montant_ttc }}"

    r_prelev = table.add_row().cells
    r_prelev[
        0
    ].text = "{% if has_prelevement %}Prélèvement ({{ taux_prelevement }}%){% endif %}"
    r_prelev[3].text = "{% if has_prelevement %}{{ prelevement }}{% endif %}"

    r_av = table.add_row().cells
    r_av[0].text = "{% if has_avance %}Avance{% endif %}"
    r_av[3].text = "{% if has_avance %}{{ avance }}{% endif %}"

    r_reg = table.add_row().cells
    r_reg[0].text = "{% if has_regule %}Régule{% endif %}"
    r_reg[3].text = "{% if has_regule %}{{ regule }}{% endif %}"

    r_opt = table.add_row().cells
    r_opt[0].text = "{% if has_autre_option %}{{ autre_option_label }}{% endif %}"
    r_opt[3].text = (
        "{% if has_autre_option %}{{ autre_option_montant }}{% endif %}"
    )

    r_net = table.add_row().cells
    r_net[0].text = "NET À PAYER"
    r_net[3].text = "{{ net_a_payer }}"

    doc.add_paragraph(
        "\nArrêté la présente facture à la somme de : {{ net_en_lettres }}"
    )

    filepath = os.path.join(TEMPLATE_DIR, "Facture_Vente_Atelier_Mobile.docx")
    doc.save(filepath)
    print(f" Template réinitialisé : {filepath}")


if __name__ == "__main__":
    generer_template_service()
    generer_template_produit()