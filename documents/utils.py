import os
import pythoncom
from django.conf import settings
from docxtpl import DocxTemplate
from num2words import num2words
from docx2pdf import convert
from docx.shared import Pt

def generer_pdf_document(document_id):
    from .models import Document
    doc_obj = Document.objects.get(id=document_id)

    # 1. Sélection dynamique du template Word
    modele_val = getattr(doc_obj, 'modele_facture', 'service')
    is_modele_produit = (modele_val == 'produit')
    
    if is_modele_produit:
        template_name = 'Facture_Vente_Atelier_Mobile.docx'
    else:
        template_name = 'Facture_Atelier_Mobile.docx'

    template_path = os.path.join(settings.BASE_DIR, 'templates', 'docx', template_name)
    if not os.path.exists(template_path):
        template_path = os.path.join(settings.BASE_DIR, 'templates', 'docx', 'Facture_Atelier_Mobile.docx')

    doc = DocxTemplate(template_path)

    # 2. Construction de la liste des lignes (compatible Jours et Quantité)
    prestations = []
    total_brut = 0

    for ligne in doc_obj.lignes.all():
        qte = ligne.quantite or 0
        pu = ligne.prix_unitaire or 0
        
        prestations.append({
            'designation': ligne.designation or '',
            'periode': ligne.periode or '',
            'taux': f"{ligne.taux_journalier:,.0f}".replace(',', ' ') if ligne.taux_journalier else '',
            'quantite': f"{qte:g}" if qte else '',
            'prix_unitaire': f"{pu:,.0f}".replace(',', ' ') if pu else '',
            'total': f"{ligne.prix_total:,.0f} FCFA".replace(',', ' ')
        })
        total_brut += ligne.prix_total

    # 3. Calculs financiers
    taux_p = doc_obj.taux_prelevement or 0
    montant_prelevement = (total_brut * taux_p) / 100 if taux_p > 0 else 0
    val_avance = doc_obj.avance or 0
    val_regule = doc_obj.regule or 0
    val_autre_montant = getattr(doc_obj, 'autre_option_montant', 0) or 0
    val_autre_label = getattr(doc_obj, 'autre_option_label', '') or ''

    net_a_payer = total_brut - montant_prelevement - val_avance - val_regule - val_autre_montant
    net_en_lettres = num2words(net_a_payer, lang='fr').capitalize() + " francs FCFA"

    # Nom de l'initiateur
    initiateur_nom = doc_obj.initiateur or doc_obj.user.get_full_name() or doc_obj.user.username

    # Formatage de la date du document
    date_doc = getattr(doc_obj, 'date_creation', None)
    date_formatted = date_doc.strftime('%d/%m/%Y') if date_doc else ''

    # Informations client : remplacement strict des None par ""
    client_obj = getattr(doc_obj, 'client', None)
    client_nom = str(client_obj) if client_obj else ''
    
    client_nif = (getattr(client_obj, 'nif', None) or getattr(doc_obj, 'client_nif', None)) or ''
    client_bp = (getattr(client_obj, 'bp', None) or getattr(doc_obj, 'client_bp', None)) or ''
    client_ville = (getattr(client_obj, 'ville', None) or getattr(doc_obj, 'client_ville', None)) or ''

    context = {
        'modele_facture': modele_val,
        'is_produit': is_modele_produit,
        'type_doc': doc_obj.get_type_document_display().upper(),
        'numero_facture': doc_obj.numero,
        'date_facture': date_formatted,
        'initiateur': initiateur_nom.upper(),
        'client_nom': client_nom,
        'client_nif': client_nif,
        'client_bp': client_bp,
        'client_ville': client_ville,
        'objet': doc_obj.objet or '',
        'prestations': prestations,
        'montant_ttc': f"{total_brut:,.0f} FCFA".replace(',', ' '),
        'has_prelevement': taux_p > 0,
        'taux_prelevement': taux_p,
        'prelevement': f"-{montant_prelevement:,.0f} FCFA".replace(',', ' '),
        'has_avance': val_avance > 0,
        'avance': f"-{val_avance:,.0f} FCFA".replace(',', ' '),
        'has_regule': val_regule > 0,
        'regule': f"-{val_regule:,.0f} FCFA".replace(',', ' '),
        'has_autre_option': val_autre_montant > 0,
        'autre_option_label': val_autre_label,
        'autre_option_montant': f"-{val_autre_montant:,.0f} FCFA".replace(',', ' '),
        'net_a_payer': f"{net_a_payer:,.0f} FCFA".replace(',', ' '),
        'net_en_lettres': net_en_lettres
    }

    temp_docx = os.path.join(settings.MEDIA_ROOT, f"temp_{doc_obj.id}.docx")
    temp_pdf = os.path.join(settings.MEDIA_ROOT, f"temp_{doc_obj.id}.pdf")

    doc.render(context)
    doc.save(temp_docx)

    # Force la police Times New Roman
    from docx import Document as DocxDocument
    doc_generated = DocxDocument(temp_docx)
    for paragraph in doc_generated.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            if not run.font.size:
                run.font.size = Pt(12)
    for table in doc_generated.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(12)
    doc_generated.save(temp_docx)

    pythoncom.CoInitialize()
    try:
        convert(temp_docx, temp_pdf)
    finally:
        pythoncom.CoUninitialize()

    return temp_pdf